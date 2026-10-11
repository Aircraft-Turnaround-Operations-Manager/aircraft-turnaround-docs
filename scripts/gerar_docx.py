"""Monta o documento final (.docx) a partir do template do professor e dos .md de especificacao/.

Uso: python gerar_docx.py TEMPLATE.docx PASTA_ESPECIFICACAO SAIDA.docx [--manter-instrucoes] [--manter-ra2]
Precisa de python-docx. O template (ESSW - Especificacao de Projeto - Template.docx) não fica no repositório.
A declaração de uso de IA vem de especificacao/00-*.md (ou de --ia=ARQUIVO) e entra antes do item 1, sem número.
Depois de gerar, abra o .docx no Word, atualize o sumário (F9) e exporte em PDF: o script não calcula páginas.
"""
import copy, io, re, struct, sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

PRODUTO = 'Aircraft Turnaround Orchestration System'
AUTORES = ['EDUARDO FABRI', 'JOÃO PEDRO CARDOSO DE LIZ', 'JOÃO VITOR CORREA OLIVEIRA', 'RODRIGO ALVES']
ANO = '2026'
LARGURA_CM, ALTURA_MAX_CM = 18.0, 22.0
ITENS = ['01-*.md', '02-*.md', '03-*.md', '04-*.md', '05-*.md', '06-*', '07-*', '08-*', '09-*.md', '10-*', '11-*.md']
INLINE = re.compile(r'(!\[[^\]]*\]\([^)]+\)|\*\*.+?\*\*|\*[^*\s](?:[^*]*[^*\s])?\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|<br\s*/?>)')
PROD = re.compile(r'^\*\*(NOME DO PRODUTO|PRODUTO):\*\*')


# ---------------------------------------------------------------- Markdown -> blocos
def ler(p):
    s = io.open(p, encoding='utf-8').read().replace('\r', '')
    return re.sub(r'<!--.*?-->', '', s, flags=re.S)


def celulas(linha):
    return [c.strip().replace('\\|', '|') for c in re.split(r'(?<!\\)\|', linha.strip().strip('|'))]


def blocos(texto, base):
    """Devolve lista de blocos: ('h', nivel, txt) ('p', txt) ('li', nivel, txt) ('img', caminho) ('t', titulos, cabecalho, linhas)."""
    out, linhas, i = [], texto.split('\n'), 0
    while i < len(linhas):
        l = linhas[i]
        s = l.strip()
        if not s:
            i += 1
        elif s.startswith('|'):
            tab = []
            while i < len(linhas) and linhas[i].strip().startswith('|'):
                tab.append(celulas(linhas[i])); i += 1
            cab = tab[0] if len(tab) > 1 and all(re.fullmatch(r':?-+:?', c) for c in tab[1]) else None
            corpo = tab[2:] if cab is not None else tab
            titulos = []  # linhas em negrito logo acima da tabela viram linhas mescladas (padrão do template)
            while out and out[-1][0] == 'p' and (PROD.match(out[-1][1]) or re.fullmatch(r'\*\*[^*]+\*\*', out[-1][1])):
                titulos.insert(0, out.pop()[1])
            out.append(('t', titulos, cab if cab and any(cab) else None, corpo))
        elif s.startswith('#'):
            n = len(s) - len(s.lstrip('#'))
            out.append(('h', n, s[n:].strip())); i += 1
        elif re.fullmatch(r'!\[[^\]]*\]\([^)]+\)', s):
            out.append(('img', base / re.search(r'\(([^)]+)\)', s).group(1))); i += 1
        elif re.match(r'(-|\d+\.)\s', s):
            out.append(('li', (len(l) - len(l.lstrip())) // 2, s)); i += 1
        else:
            out.append(('p', s)); i += 1
    return out


# ---------------------------------------------------------------- escrita no docx
class Montador:
    def __init__(self, doc):
        self.doc, self.citadas = doc, set()
        self.tem = {s.name for s in doc.styles}

    def _mover(self, el, ancora):
        ancora.addprevious(el)

    def inline(self, par, txt, base=None, negrito=False, italico=False, tam=None, larg_img=LARGURA_CM):
        self.citadas |= {int(n) for n in re.findall(r'\[(\d+)\]', txt)}
        for parte in INLINE.split(txt):
            if not parte:
                continue
            if re.fullmatch(r'<br\s*/?>', parte):
                par.add_run().add_break()
            elif parte.startswith('!['):
                self.figura(par, (base or Path('.')) / re.search(r'\(([^)]+)\)', parte).group(1), larg_img, 12.0)
            elif parte.startswith('**'):
                self.inline(par, parte[2:-2], base, True, italico, tam, larg_img)
            elif parte.startswith('`'):
                r = par.add_run(parte[1:-1]); r.bold, r.italic = negrito or None, italico or None
                if tam: r.font.size = Pt(tam)
            elif parte.startswith('['):
                self.inline(par, re.match(r'\[([^\]]+)\]', parte).group(1), base, negrito, italico, tam, larg_img)
            elif parte.startswith('*') and parte.endswith('*') and len(parte) > 2:
                self.inline(par, parte[1:-1], base, negrito, True, tam, larg_img)
            else:
                r = par.add_run(parte); r.bold, r.italic = negrito or None, italico or None
                if tam: r.font.size = Pt(tam)

    def figura(self, par, caminho, larg_cm, alt_cm):
        if not Path(caminho).exists():
            par.add_run('[imagem não encontrada: %s]' % caminho); return
        d = open(caminho, 'rb').read(32)
        w, h = struct.unpack('>II', d[16:24]) if d[1:4] == b'PNG' else (1, 1)
        larg = min(larg_cm, alt_cm * w / h)
        par.add_run().add_picture(str(caminho), width=Cm(larg))

    def paragrafo(self, ancora, txt, base=None, estilo=None, **kw):
        p = self.doc.add_paragraph(style=estilo if estilo in self.tem else None)
        self.inline(p, txt, base, **kw)
        self._mover(p._p, ancora); return p

    def titulo(self, ancora, nivel, txt, sem_num=False):
        est = 'Heading %d' % nivel
        if not sem_num and nivel >= 2:  # o estilo já numera: tira o número escrito no Markdown
            txt = re.sub(r'^\d+(\.\d+)+\s+', '', txt)
        if est in self.tem:
            p = self.doc.add_paragraph(style=est); p.add_run(txt)
            if sem_num:
                numPr = OxmlElement('w:numPr'); nid = OxmlElement('w:numId'); nid.set(qn('w:val'), '0')
                numPr.append(nid); p._p.get_or_add_pPr().append(numPr)
        else:
            p = self.doc.add_paragraph(); r = p.add_run(txt); r.bold = True
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(10)
        self._mover(p._p, ancora)

    def imagem(self, ancora, caminho):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.figura(p, caminho, LARGURA_CM, ALTURA_MAX_CM)
        self._mover(p._p, ancora)

    def lista(self, ancora, nivel, txt, base):
        m = re.match(r'(-|\d+\.)\s+(.*)', txt)
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6 + 0.6 * nivel)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        self.inline(p, ('• ' if m.group(1) == '-' else m.group(1) + ' ') + m.group(2), base)
        self._mover(p._p, ancora)

    def _sombra(self, cel, cor='D9D9D9'):
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), cor)
        cel._tc.get_or_add_tcPr().append(shd)

    def _cel(self, cel, txt, base, negrito=False, tam=9, larg_img=11.0):
        cel.text = ''
        self.inline(cel.paragraphs[0], txt, base, negrito=negrito, tam=tam, larg_img=larg_img)

    def tabela(self, ancora, titulos, cab, linhas, base, tam=9):
        ncol = max([len(cab or [])] + [len(l) for l in linhas] + [1])
        t = self.doc.add_table(rows=0, cols=ncol)
        t.style = 'Table Grid' if 'Table Grid' in self.tem else t.style
        t.autofit = True
        for tit in titulos:
            c = t.add_row().cells
            m = c[0].merge(c[-1]) if ncol > 1 else c[0]
            self._cel(m, tit, base, negrito=not PROD.match(tit), tam=tam + 1); self._sombra(m)
            if not PROD.match(tit): m.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if cab:
            c = t.add_row().cells
            for k in range(ncol):
                self._cel(c[k], cab[k] if k < len(cab) else '', base, negrito=True, tam=tam); self._sombra(c[k])
            trPr = t.rows[-1]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
        for l in linhas:
            c = t.add_row().cells
            for k in range(ncol):
                self._cel(c[k], l[k] if k < len(l) else '', base, tam=tam, larg_img=11.0 if ncol <= 2 else 6.0)
        self._mover(t._tbl, ancora)
        self._mover(self.doc.add_paragraph()._p, ancora)

    def escrever(self, ancora, bls, base, tam=9, h1=False, sem_num=False):
        for b in bls:
            if b[0] == 'h':
                if b[1] == 1 and not h1: continue
                self.titulo(ancora, b[1], b[2], sem_num=(sem_num or b[1] == 1))
            elif b[0] == 'p':
                self.paragrafo(ancora, b[1], base)
            elif b[0] == 'li':
                self.lista(ancora, b[1], b[2], base)
            elif b[0] == 'img':
                self.imagem(ancora, b[1])
            elif b[0] == 't':
                self.tabela(ancora, b[1], b[2], b[3], base, tam)


# ---------------------------------------------------------------- regras por item
def fontes_unidas(bls):
    """Junta as linhas 'Fonte(s) citada(s)' de vários arquivos em uma só."""
    nums, resto = set(), []
    for b in bls:
        if b[0] == 'p' and re.match(r'Fontes? citadas?:', b[1]):
            nums |= {int(n) for n in re.findall(r'\[(\d+)\]', b[1])}
        else:
            resto.append(b)
    if nums:
        ns = ['[%d]' % n for n in sorted(nums)]
        lista = ns[0] if len(ns) == 1 else ', '.join(ns[:-1]) + ' e ' + ns[-1]
        resto.append(('p', ('Fonte citada: ' if len(ns) == 1 else 'Fontes citadas: ') + lista + ', conforme a numeração de `pesquisa/fontes.md`.'))
    return resto


def tabela_unica(bls):
    """Itens 6 e 8: as tabelas das áreas, de mesmo cabeçalho, viram uma tabela só."""
    out, alvo = [], None
    for b in bls:
        if b[0] == 't' and b[2] and alvo is not None and alvo[2] == b[2]:
            alvo[3].extend(b[3])
        else:
            if b[0] == 't' and b[2] and alvo is None:
                b = ('t', list(b[1]), b[2], list(b[3])); alvo = b
            out.append(b)
    if alvo is not None and not alvo[1]:  # a linha PRODUTO do 00-item sobe para o topo da tabela
        for k, b in enumerate(out):
            if b[0] == 'p' and PROD.match(b[1]):
                alvo[1].append(out.pop(k)[1]); break
    return fontes_unidas(out)


def estorias(bls):
    """Item 7: cada estória vira uma tabela no formato do template."""
    out, i = [], 0
    while i < len(bls):
        b = bls[i]
        if b[0] == 'h' and b[1] == 2 and re.match(r'US', b[2]):
            campos, crit, j = [], [], i + 1
            while j < len(bls) and not (bls[j][0] == 'h' and bls[j][1] <= 2):
                x = bls[j]
                if x[0] == 'p' and re.match(r'\*\*(COMO|POSSO|PARA):\*\*', x[1]): campos.append(x[1])
                elif x[0] == 't': crit = x[3]
                j += 1
            out.append(('t', ['**%s**' % b[2].replace('*', '')], None, [['<br>'.join(campos)], ['**Critérios de Aceite:**']] + crit))
            i = j
        else:
            out.append(b); i += 1
    return fontes_unidas(out)


def carregar(pasta, padrao):
    alvo = sorted(pasta.glob(padrao))[0]
    arqs = [alvo] if alvo.is_file() else sorted(alvo.glob('*.md'))
    bls = []
    for a in arqs:
        t = ler(a)
        if alvo.is_dir() and a.name != '00-item.md':  # tabela de trabalho do item 9, no início dos arquivos do item 10
            t = re.sub(r'(?s)\*\*Casos de uso da área.*?(?=^## )', '', t, count=1, flags=re.M)
        bls += blocos(t, a.parent)
    n = padrao[:2]
    if n in ('06', '08'): bls = tabela_unica(bls)
    if n == '07': bls = estorias(bls)
    if n == '10': bls = fontes_unidas(bls)
    return bls


# ---------------------------------------------------------------- template
def trocar_texto(par, novo):
    rs = par.runs
    if rs:
        rs[0].text = novo
        for r in rs[1:]: r.text = ''
    else:
        par.add_run(novo)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    manter_instr, manter_ra2 = '--manter-instrucoes' in sys.argv, '--manter-ra2' in sys.argv
    template, pasta, saida = Path(args[0]), Path(args[1]), Path(args[2])
    doc = Document(str(template)); m = Montador(doc)
    # capa e rodapé
    autores = iter(AUTORES)
    for p in doc.paragraphs:
        t = p.text.strip()
        if re.fullmatch(r'NOME AUTOR \d', t): trocar_texto(p, next(autores))
        elif 'NOME DO PRODUTO DE SOFTWARE' in t: trocar_texto(p, '- %s -' % PRODUTO.upper())
        elif t == '2025': trocar_texto(p, ANO)
    for s in doc.sections:
        for p in s.footer.paragraphs:
            for r in p.runs:
                if 'Nome do Produto de Software' in r.text: r.text = r.text.replace('Nome do Produto de Software', PRODUTO)
    MC = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
    for ac in list(doc.element.body.iter(MC + 'AlternateContent')):
        if 'Aviso' in ''.join(t.text or '' for t in ac.iter(qn('w:t'))):
            r = ac.getparent(); r.getparent().remove(r)
    for parte in [doc.element.body] + [s.footer._element for s in doc.sections] + [s.header._element for s in doc.sections]:
        for c in parte.iter(qn('w:color')):
            if c.get(qn('w:val')) == '00B0F0': c.set(qn('w:val'), '000000')
    for it in doc.element.body.iter(qn('w:instrText')):
        if it.text and 'TOC' in it.text: it.text = it.text.replace('"1-3"', '"1-2"')
    # seções do template
    corpo = doc.element.body
    h1 = [e for e in corpo.iterchildren(qn('w:p')) if Paragraph(e, doc).style.name == 'Heading 1']
    assert len(h1) >= 11, len(h1)
    ia = next((Path(a.split('=', 1)[1]) for a in sys.argv if a.startswith('--ia=')), None) or next(iter(sorted(pasta.glob('00-*.md'))), None)
    if ia and ia.exists():
        m.escrever(h1[0], blocos(ler(ia), ia.parent), ia.parent, h1=True)
        p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE); h1[0].addprevious(p._p)
    sect_final = corpo.find(qn('w:sectPr'))
    eh_pb = lambda e: e.tag == qn('w:p') and any(b.get(qn('w:type')) == 'page' for b in e.iter(qn('w:br')))
    for k, cab in enumerate(h1):
        fim = h1[k + 1] if k + 1 < len(h1) else sect_final
        meio = []
        e = cab.getnext()
        while e is not None and e is not fim:
            meio.append(e); e = e.getnext()
        if k >= 11:  # itens 12 a 16 são do RA2
            if not manter_ra2:
                for x in meio + [cab]: corpo.remove(x)
            continue
        pb = next((x for x in reversed(meio) if eh_pb(x)), None)
        viu_tabela = False
        for x in meio:
            if x is pb: continue
            viu_tabela = viu_tabela or x.tag == qn('w:tbl')
            if not manter_instr or viu_tabela or not Paragraph(x, doc).text.strip() if x.tag == qn('w:p') else True:
                corpo.remove(x)
        ancora = pb if pb is not None else fim
        m.escrever(ancora, carregar(pasta, ITENS[k]), pasta)
    # apêndice A e referências, no fim
    ancora = corpo.find(qn('w:sectPr'))
    ant = ancora.getprevious()
    if ant is None or not eh_pb(ant):
        p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE); ancora.addprevious(p._p)
    ap = pasta / 'apendice-a-matriz-de-rastreabilidade.md'
    if ap.exists():
        m.escrever(ancora, blocos(ler(ap), pasta), pasta, tam=7, h1=True, sem_num=True)
    fontes = pasta.parent / 'pesquisa' / 'fontes.md'
    if fontes.exists() and m.citadas:
        p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE); ancora.addprevious(p._p)
        m.titulo(ancora, 1, 'REFERÊNCIAS', sem_num=True)
        cit = set(m.citadas)
        for l in ler(fontes).split('\n'):
            mm = re.match(r'(\d+)\.\s+(.*)', l.strip())
            if mm and int(mm.group(1)) in cit:
                m.paragrafo(ancora, '[%s] %s' % (mm.group(1), mm.group(2)), tam=9)
    laranja = sum(1 for c in doc.element.body.iter(qn('w:color')) if c.get(qn('w:val')) == 'ED7D31')
    doc.save(str(saida))
    print('textos em laranja restantes:', laranja)
    print('gerado', saida, '| fontes citadas', len(m.citadas))


if __name__ == '__main__':
    main()
